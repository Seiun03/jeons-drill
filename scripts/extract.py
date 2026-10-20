#!/usr/bin/env python3
"""강의자료(PDF/PPTX)를 페이지·슬라이드 단위 마크다운으로 변환한다.

사용법 (과목 폴더에서):
    python <SKILL_DIR>/scripts/extract.py            # materials/ -> .drill/notes/
    python <SKILL_DIR>/scripts/extract.py --force    # 이미 변환된 파일도 다시 변환

출력:
    .drill/notes/<파일명>.md      ## [p.N] 또는 ## [slide N] 단위로 구분된 텍스트
    .drill/images/...             그림이 중요해 보이는 페이지/슬라이드 이미지
    .drill/extract_report.json    파일별 페이지 수, 이미지 검토가 필요한 위치 목록

이미지 검토가 필요한 위치에는 노트에 다음 표시가 들어간다:
    <!-- IMAGE_REVIEW: .drill/images/xxx.png -->
에이전트는 그 이미지를 직접 보고, 표시 바로 아래에 그림 내용을 글로 적어 넣어야 한다.

필요 패키지: pip install pymupdf python-pptx
.ppt(구형 형식)는 LibreOffice(soffice)가 있으면 자동으로 .pptx로 변환한다.
"""
import argparse
import json
import shutil
import subprocess
import sys
from pathlib import Path

try:
    sys.stdout.reconfigure(encoding="utf-8")
except Exception:
    pass

LOW_TEXT_CHARS = 80        # 이보다 글자가 적으면 그림 위주 페이지로 판단
IMAGE_AREA_RATIO = 0.25    # 그림이 페이지 면적의 이 비율 이상이면 검토 대상
RENDER_DPI = 110


def need(module, pip_name):
    try:
        return __import__(module)
    except ImportError:
        print(f"[오류] {pip_name} 패키지가 필요합니다: pip install {pip_name}")
        sys.exit(1)


# ---------------------------------------------------------------- PDF
def extract_pdf(path: Path, notes_dir: Path, img_dir: Path):
    pymupdf = need("pymupdf", "pymupdf")
    doc = pymupdf.open(path)
    lines = [f"# {path.stem}", f"<!-- source: {path.name}, pages: {doc.page_count} -->", ""]
    flagged = []
    low_text_pages = 0
    for i, page in enumerate(doc, start=1):
        text = page.get_text("text").strip()
        page_area = abs(page.rect) or 1
        img_area = 0.0
        try:
            for info in page.get_image_info():
                x0, y0, x1, y1 = info["bbox"]
                img_area += max(0, x1 - x0) * max(0, y1 - y0)
        except Exception:
            pass
        has_drawings = False
        try:
            has_drawings = len(page.get_drawings()) > 40  # 도표·그래프가 벡터로 그려진 경우
        except Exception:
            pass
        low_text = len(text) < LOW_TEXT_CHARS
        if low_text:
            low_text_pages += 1
        review = low_text or (img_area / page_area) >= IMAGE_AREA_RATIO or has_drawings

        lines.append(f"## [p.{i}]")
        lines.append(text if text else "(텍스트 없음)")
        if review:
            out = img_dir / f"{path.stem}_p{i}.png"
            page.get_pixmap(dpi=RENDER_DPI).save(out)
            rel = out.as_posix()
            lines.append(f"<!-- IMAGE_REVIEW: {rel} -->")
            flagged.append({"loc": f"p.{i}", "image": rel})
        lines.append("")
    count = doc.page_count
    doc.close()
    (notes_dir / f"{path.stem}.md").write_text("\n".join(lines), encoding="utf-8")
    scanned = count > 0 and low_text_pages / count > 0.7
    return {"pages": count, "flagged": flagged, "scanned_suspect": scanned}


# ---------------------------------------------------------------- PPTX
def shape_texts(shape, out):
    from pptx.enum.shapes import MSO_SHAPE_TYPE
    if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
        for s in shape.shapes:
            shape_texts(s, out)
        return
    if getattr(shape, "has_table", False) and shape.has_table:
        for row in shape.table.rows:
            out.append("| " + " | ".join(c.text.strip().replace("\n", " ") for c in row.cells) + " |")
        return
    if getattr(shape, "has_text_frame", False) and shape.has_text_frame:
        for p in shape.text_frame.paragraphs:
            t = "".join(r.text for r in p.runs).strip()
            if t:
                out.append("  " * p.level + "- " + t)


def collect_pictures(shapes, out):
    from pptx.enum.shapes import MSO_SHAPE_TYPE
    for s in shapes:
        if s.shape_type == MSO_SHAPE_TYPE.GROUP:
            collect_pictures(s.shapes, out)
        elif s.shape_type == MSO_SHAPE_TYPE.PICTURE:
            out.append(s)
        elif getattr(s, "has_chart", False) and s.has_chart:
            out.append(s)


def extract_pptx(path: Path, notes_dir: Path, img_dir: Path):
    need("pptx", "python-pptx")
    from pptx import Presentation
    prs = Presentation(path)
    slide_area = (prs.slide_width or 1) * (prs.slide_height or 1)
    lines = [f"# {path.stem}", f"<!-- source: {path.name}, slides: {len(prs.slides)} -->", ""]
    flagged = []
    for i, slide in enumerate(prs.slides, start=1):
        texts = []
        for shape in slide.shapes:
            shape_texts(shape, texts)
        body = "\n".join(texts)
        lines.append(f"## [slide {i}]")
        lines.append(body if body else "(텍스트 없음)")
        if slide.has_notes_slide:
            note = slide.notes_slide.notes_text_frame.text.strip()
            if note:
                lines.append("")
                lines.append("> 발표자 노트: " + note.replace("\n", "\n> "))

        pics = []
        collect_pictures(slide.shapes, pics)
        big = [p for p in pics if p.width and p.height and (p.width * p.height) / slide_area >= 0.08]
        if big or (len(body) < LOW_TEXT_CHARS and pics):
            saved = []
            for k, p in enumerate(big or pics, start=1):
                if not hasattr(p, "image"):
                    lines.append("<!-- IMAGE_REVIEW: 차트 개체 (원본 슬라이드 확인 필요) -->")
                    flagged.append({"loc": f"slide {i}", "image": None, "note": "chart"})
                    continue
                ext = p.image.ext or "png"
                out = img_dir / f"{path.stem}_s{i}_{k}.{ext}"
                out.write_bytes(p.image.blob)
                saved.append(out.as_posix())
            for rel in saved:
                lines.append(f"<!-- IMAGE_REVIEW: {rel} -->")
                flagged.append({"loc": f"slide {i}", "image": rel})
        lines.append("")
    (notes_dir / f"{path.stem}.md").write_text("\n".join(lines), encoding="utf-8")
    return {"pages": len(prs.slides), "flagged": flagged, "scanned_suspect": False}


# ---------------------------------------------------------------- .ppt 변환
def convert_ppt(path: Path, work: Path):
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if not soffice:
        return None
    work.mkdir(parents=True, exist_ok=True)
    subprocess.run([soffice, "--headless", "--convert-to", "pptx", "--outdir", str(work), str(path)],
                   capture_output=True, timeout=300)
    out = work / (path.stem + ".pptx")
    return out if out.exists() else None


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--root", default=".", help="과목 폴더 (기본: 현재 폴더)")
    ap.add_argument("--force", action="store_true", help="이미 변환된 파일도 다시 변환")
    args = ap.parse_args()

    root = Path(args.root)
    materials = root / "materials"
    drill = root / ".drill"
    notes_dir, img_dir = drill / "notes", drill / "images"
    if not materials.is_dir():
        print(f"[오류] {materials} 폴더가 없습니다. 강의자료를 materials/ 에 넣어 주세요.")
        sys.exit(1)
    notes_dir.mkdir(parents=True, exist_ok=True)
    img_dir.mkdir(parents=True, exist_ok=True)

    report_path = drill / "extract_report.json"
    report = json.loads(report_path.read_text(encoding="utf-8")) if report_path.exists() else {}

    files = sorted(p for p in materials.iterdir() if p.is_file() and not p.name.startswith((".", "~$")))
    for f in files:
        ext = f.suffix.lower()
        note = notes_dir / f"{f.stem}.md"
        if note.exists() and not args.force:
            print(f"[건너뜀] {f.name} (이미 변환됨, 다시 하려면 --force)")
            continue
        try:
            if ext == ".pdf":
                r = extract_pdf(f, notes_dir, img_dir)
            elif ext == ".pptx":
                r = extract_pptx(f, notes_dir, img_dir)
            elif ext == ".ppt":
                conv = convert_ppt(f, drill / "converted")
                if not conv:
                    print(f"[실패] {f.name}: .ppt는 PowerPoint에서 .pptx로 저장해 주세요 (LibreOffice 없음)")
                    continue
                r = extract_pptx(conv, notes_dir, img_dir)
            else:
                print(f"[무시] {f.name}: 지원하지 않는 형식")
                continue
        except Exception as e:
            print(f"[실패] {f.name}: {e}")
            continue
        report[f.name] = {"note": note.as_posix(), **r}
        flag = f", 이미지 검토 {len(r['flagged'])}곳" if r["flagged"] else ""
        scan = "  ※ 스캔본 의심: 대부분 페이지에 텍스트가 없음" if r["scanned_suspect"] else ""
        unit = "쪽" if ext == ".pdf" else "장"
        print(f"[완료] {f.name}: {r['pages']}{unit}{flag}{scan}")

    report_path.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    total_flag = sum(len(v.get("flagged", [])) for v in report.values())
    print(f"\n노트: {notes_dir.as_posix()}  /  이미지 검토 필요: 총 {total_flag}곳  /  보고서: {report_path.as_posix()}")


if __name__ == "__main__":
    main()
